import datetime

from openpyxl import load_workbook

from src.judge import judge_record
from src.models import OrderRecord, ProcessStep
from src.process_history import ProcessHistory
from src.process_master import ProcessMaster
from src.report_excel import write_excel_report


def make_judged_record(**overrides) -> OrderRecord:
    defaults = dict(
        order_no="PO-000001",
        drawing_no="312127272601",
        product_name="SPB M10x13T L",
        qty=1,
        order_date=datetime.date(2026, 8, 1),
        company_deadline=datetime.date(2026, 9, 1),
        customer_deadline=datetime.date(2026, 9, 5),
        remaining_business_days=20,
    )
    defaults.update(overrides)
    record = OrderRecord(**defaults)
    judge_record(record)
    return record


def test_write_excel_report_creates_expected_sheets(tmp_path):
    records = [
        make_judged_record(order_no="ORD-DELAYED", remaining_business_days=-3),
        make_judged_record(order_no="ORD-RISK", remaining_business_days=2),
        make_judged_record(order_no="ORD-UNDETERMINED", company_deadline=None, remaining_business_days=None),
    ]
    output_path = tmp_path / "report.xlsx"

    write_excel_report(records, output_path, generated_at=datetime.date(2026, 8, 3), process_master=ProcessMaster())

    assert output_path.exists()
    wb = load_workbook(output_path)
    assert wb.sheetnames == [
        "メインサマリー",
        "①納期遅延リスト",
        "②納期遅延リスクリスト",
        "判定不能・要確認",
        "工程別仕掛中ランキング",
    ]
    assert wb["①納期遅延リスト"]["A2"].value == "ORD-DELAYED"
    assert wb["②納期遅延リスクリスト"]["A2"].value == "ORD-RISK"
    assert wb["判定不能・要確認"]["A2"].value == "ORD-UNDETERMINED"


def test_write_excel_report_adds_process_history_sheets(tmp_path):
    record = make_judged_record(
        order_no="ORD-RISK", remaining_business_days=2,
        processes=[ProcessStep(2, "HB", "F1", "d1", "d2", "オーダー確定前")],
    )

    process_history = ProcessHistory()
    process_history._actual_durations["HB"] = [10]
    process_history._weekly_load_by_department[(datetime.date(2026, 8, 31), "1223")] = 12.5
    process_history._weekly_load_by_machine[(datetime.date(2026, 8, 31), "F1")] = 12.5

    output_path = tmp_path / "report.xlsx"
    write_excel_report(
        [record], output_path, generated_at=datetime.date(2026, 9, 8),
        process_master=ProcessMaster(), process_history=process_history,
    )

    wb = load_workbook(output_path)
    assert "実績ベース納期充足予測" in wb.sheetnames
    assert "週別負荷(部署別)" in wb.sheetnames
    assert "週別負荷(設備別)" in wb.sheetnames
    assert wb["実績ベース納期充足予測"]["A2"].value == "ORD-RISK"
    assert wb["週別負荷(部署別)"]["B2"].value == 12.5
    assert wb["週別負荷(設備別)"]["C2"].value == 12.5
