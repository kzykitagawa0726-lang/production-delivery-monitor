"""仕入(外注)累積データから、図番・工程コードごとの仕入先実績を集計する。

kii-san提供の複数年分の仕入データExcel(基幹システムからのエクスポート)を
読み込み、「この図番/工程コードなら過去にどの仕入先が対応してきたか」を
集計する。①②で停滞している工程の「代替先候補」提示に使う(2026-09-08 設計合意)。

- 図番での紐付けが最優先(実データで78.2%が週次受注データの図番と一致することを確認済み)。
  一致がなければ工程コード単位の一般的な実績にフォールバックする。
- 仕入単価・金額はセンシティブな財務情報のため、集計結果には一切含めない
  (件数・最終利用日のみ)。
- 工程コードが空欄の行(材料そのものの仕入と推測、kii-san確認済み)は集計対象外。
- 入力ファイルは受注データと異なり更新頻度が低く、かつ合計で数十万行規模になるため、
  集計結果のみをローカルキャッシュ(.cache/supplier_history.json、gitには含めない)に
  保存し、入力ファイル一式が変わらない限り再読み込みを省略する。
"""
from __future__ import annotations

import datetime
import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from src.process_master import normalize_process_code

DEFAULT_CACHE_PATH = Path(".cache/supplier_history.json")
# 集計ロジックを変更したら上げる。ソースファイルが同じでもキャッシュを無効化し、
# 古いロジックで集計された結果を誤って使い続けないようにするためのガード。
CACHE_SCHEMA_VERSION = 1

# 2022年のみ列数が少ない簡易フォーマット。2023年以降は55列共通フォーマット。
HEADER_19 = [
    "得意先ＣＤ", "仕入先ＣＤ", "担当CD", "商品名", "数量", "仕入単価", "金額（支払額）",
    "受注№", "行№", "情報ＮＯ", "注文№", "図番", "材質", "寸法", "伝票日付", "工程コード",
    "明細１", "明細２", "明細４",
]
HEADER_55 = [
    "部門ＣＤ", "得意先ＣＤ", "仕入先ＣＤ", "伝票日付(未)", "伝票種類", "伝票№", "行", "取引区分",
    "FILLER", "担当CD", "請求発行区分", "台帳発行区分", "翌月区分", "翌締区分", "商品グループ",
    "科目", "予備", "商品名", "単位", "数量", "仕入単価", "金額（支払額）", "予備２", "備考",
    "受注№", "行№", "情報ＮＯ", "注文№", "予備３", "完了区分", "図番", "貴社注文番号", "製品区分",
    "材質", "寸法", "仕入区分", "予備", "伝票日付", "予備", "工程コード", "予備", "更新", "抽出",
    "受注", "明細１", "明細２", "明細３", "明細４", "明細５", "明細６", "窓口ＣＤ", "消費税",
    "予備", "予備", "予備",
]


@dataclass
class SupplierSuggestion:
    supplier_code: str
    process_code: Optional[str]  # 図番一致の場合、最頻の工程コード。工程コード一致の場合はクエリと同じ
    match_type: str  # "drawing"(図番一致) or "process_code"(工程コード一致)
    count: int
    last_used: Optional[datetime.date]

    @property
    def label(self) -> str:
        kind = "図番一致" if self.match_type == "drawing" else "工程一致"
        last = self.last_used.isoformat() if self.last_used else "不明"
        return f"仕入先{self.supplier_code}({kind} {self.count}回, 最終{last})"


def _parse_date(value) -> Optional[datetime.date]:
    if not isinstance(value, int):
        return None
    text = str(value)
    if len(text) != 8:
        return None
    try:
        return datetime.date(int(text[0:4]), int(text[4:6]), int(text[6:8]))
    except ValueError:
        return None


def _iter_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header = HEADER_19 if ws.max_column <= 20 else HEADER_55
    idx = {name: i for i, name in enumerate(header)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield idx, row


class SupplierHistory:
    def __init__(self) -> None:
        # by_drawing[図番][仕入先CD] = {"count": int, "last_used": date|None, "process_codes": Counter}
        self._by_drawing: dict[str, dict[str, dict]] = defaultdict(lambda: defaultdict(
            lambda: {"count": 0, "last_used": None, "process_codes": Counter()}
        ))
        # by_process[工程コード][仕入先CD] = {"count": int, "last_used": date|None}
        self._by_process: dict[str, dict[str, dict]] = defaultdict(lambda: defaultdict(
            lambda: {"count": 0, "last_used": None}
        ))

    @classmethod
    def load(cls, paths: list[Path], cache_path: Path = DEFAULT_CACHE_PATH) -> "SupplierHistory":
        signature = {"__schema__": CACHE_SCHEMA_VERSION, **{str(p): p.stat().st_mtime for p in paths}}
        cached = _try_load_cache(cache_path, signature)
        if cached is not None:
            return cached

        history = cls()
        for path in paths:
            history._ingest_file(path)
        _save_cache(cache_path, signature, history)
        return history

    def _ingest_file(self, path: Path) -> None:
        for idx, row in _iter_rows(path):
            raw_code = row[idx["工程コード"]]
            if raw_code in (None, "", " "):
                continue  # 工程コード空欄=材料そのものの仕入。集計対象外(kii-san確認済み)
            process_code = normalize_process_code(raw_code)

            supplier_code = row[idx["仕入先ＣＤ"]]
            if supplier_code in (None, ""):
                continue
            supplier_code = str(supplier_code).strip()

            drawing_no_raw = row[idx["図番"]]
            drawing_no = str(drawing_no_raw).strip() if drawing_no_raw not in (None, "") else None

            date = _parse_date(row[idx["伝票日付"]])

            proc_entry = self._by_process[process_code][supplier_code]
            proc_entry["count"] += 1
            if date and (proc_entry["last_used"] is None or date > proc_entry["last_used"]):
                proc_entry["last_used"] = date

            if drawing_no:
                dwg_entry = self._by_drawing[drawing_no][supplier_code]
                dwg_entry["count"] += 1
                dwg_entry["process_codes"][process_code] += 1
                if date and (dwg_entry["last_used"] is None or date > dwg_entry["last_used"]):
                    dwg_entry["last_used"] = date

    def suggest(self, drawing_no: Optional[str], process_code: Optional[str], top_n: int = 3) -> list[SupplierSuggestion]:
        if drawing_no and drawing_no in self._by_drawing:
            entries = self._by_drawing[drawing_no]
            ranked = sorted(entries.items(), key=lambda kv: kv[1]["count"], reverse=True)[:top_n]
            return [
                SupplierSuggestion(
                    supplier_code=supplier,
                    process_code=data["process_codes"].most_common(1)[0][0] if data["process_codes"] else None,
                    match_type="drawing",
                    count=data["count"],
                    last_used=data["last_used"],
                )
                for supplier, data in ranked
            ]

        if process_code:
            normalized = normalize_process_code(process_code)
            if normalized in self._by_process:
                entries = self._by_process[normalized]
                ranked = sorted(entries.items(), key=lambda kv: kv[1]["count"], reverse=True)[:top_n]
                return [
                    SupplierSuggestion(
                        supplier_code=supplier,
                        process_code=normalized,
                        match_type="process_code",
                        count=data["count"],
                        last_used=data["last_used"],
                    )
                    for supplier, data in ranked
                ]

        return []

    def to_cache_dict(self) -> dict:
        return {
            "by_drawing": {
                drawing: {
                    supplier: {
                        "count": data["count"],
                        "last_used": data["last_used"].isoformat() if data["last_used"] else None,
                        "process_codes": dict(data["process_codes"]),
                    }
                    for supplier, data in suppliers.items()
                }
                for drawing, suppliers in self._by_drawing.items()
            },
            "by_process": {
                code: {
                    supplier: {
                        "count": data["count"],
                        "last_used": data["last_used"].isoformat() if data["last_used"] else None,
                    }
                    for supplier, data in suppliers.items()
                }
                for code, suppliers in self._by_process.items()
            },
        }

    @classmethod
    def from_cache_dict(cls, payload: dict) -> "SupplierHistory":
        history = cls()
        for drawing, suppliers in payload["by_drawing"].items():
            for supplier, data in suppliers.items():
                entry = history._by_drawing[drawing][supplier]
                entry["count"] = data["count"]
                entry["last_used"] = datetime.date.fromisoformat(data["last_used"]) if data["last_used"] else None
                entry["process_codes"] = Counter(data["process_codes"])
        for code, suppliers in payload["by_process"].items():
            for supplier, data in suppliers.items():
                entry = history._by_process[code][supplier]
                entry["count"] = data["count"]
                entry["last_used"] = datetime.date.fromisoformat(data["last_used"]) if data["last_used"] else None
        return history


def _try_load_cache(cache_path: Path, signature: dict) -> Optional[SupplierHistory]:
    if not cache_path.exists():
        return None
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    if payload.get("signature") != signature:
        return None
    return SupplierHistory.from_cache_dict(payload["data"])


def _save_cache(cache_path: Path, signature: dict, history: SupplierHistory) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"signature": signature, "data": history.to_cache_dict()}
    cache_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
