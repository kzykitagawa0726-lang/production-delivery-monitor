"""基幹システムからエクスポートされた受注データExcelの読み込み。

実データ確認の結果(2026-09-08 kii-san確認)、列構成は以下の通り(複合ヘッダー、
実質的な列名は「図番/品名」が現れる行にある。通常は5行目、データはその直後から)。

  A: 図番/品名(改行区切り)
  B: 製造オーダー№
  C: 客先注番
  D: 数量
  E: 受注日\n自社納期\n顧客納期(改行区切り。自社納期は全件実データ、
     受注日・顧客納期は「00年01月00日」が未設定プレースホルダーとして
     入ることがある。内示・先行手配案件と推測されるが①②の判定からは
     除外しない方針、とkii-san確認済み)
  F: 残日(基幹システム側で計算済みの残営業日数。そのまま使用)
  G: 得意先NO
  H: 停滞日(数値、または前工程の実績が無い場合は文字列「前工程実績無し」)
  I: 製番(旧基幹システムキー。空欄は新基幹システムで手配した案件を意味し正常)
  J: MEMO1(営業が客先とすり合わせた納期。8桁のYYYYMMDD数値の場合は日付として、
     それ以外はテキストメモとして扱う)
  K: MEMO2(営業・生産管理の自由記載。0は「記載なし」として扱う)
  L: MEMO3(月次生産計画表の売上区分。週次の出荷予定バケット)

M列以降(工程進捗、最大20スロット)は以下の形式:
  スロット1(M列): 固定文言(通常「材料入荷済み」)
  スロット2〜20(N列〜): "工程コード\n加工先\n日付1\n日付2\nステータス" の5行形式
    ステータスは 作業完了 / オーダー確定 / オーダー確定前 / 発注中 の4種のみ確認。
    工程コード→カテゴリの対応表(PDF由来)が届き次第、process_master.py で
    カテゴリ分類・ボトルネック強調を有効化する。
"""
from __future__ import annotations

import datetime
import re
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from src.models import OrderRecord, ProcessStep

COL_DRAWING_PRODUCT = 0  # A
COL_ORDER_NO = 1  # B
COL_CUSTOMER_ORDER_NO = 2  # C
COL_QTY = 3  # D
COL_DEADLINE_RAW = 4  # E
COL_REMAINING_DAYS = 5  # F
COL_CUSTOMER_NO = 6  # G
COL_STAGNATION = 7  # H
COL_OLD_SYSTEM_NO = 8  # I
COL_MEMO1 = 9  # J
COL_MEMO2 = 10  # K
COL_MEMO3 = 11  # L

PROCESS_SLOT_START_COL = 12  # M列(スロット1, 材料入荷状況)
PROCESS_SLOT_COUNT = 20

HEADER_MARKER = "図番/品名"
DEADLINE_PLACEHOLDER = "00年01月00日"

DEADLINE_LINE_PATTERN = re.compile(r"(\d{2})年(\d{1,2})月(\d{1,2})日")


def _clean_str(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_deadline_line(line: str) -> Optional[datetime.date]:
    line = line.strip()
    if not line or line == DEADLINE_PLACEHOLDER:
        return None
    m = DEADLINE_LINE_PATTERN.match(line)
    if not m:
        return None
    yy, mm, dd = m.groups()
    try:
        return datetime.date(2000 + int(yy), int(mm), int(dd))
    except ValueError:
        return None


def _parse_qty(value) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _parse_remaining_business_days(value) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _split_drawing_product(value) -> tuple[str, Optional[str]]:
    text = "" if value is None else str(value)
    parts = text.split("\n", 1)
    drawing_no = parts[0].strip()
    product_name = parts[1].strip() if len(parts) > 1 else None
    return drawing_no, (product_name or None)


def _parse_deadline_triplet(value) -> tuple[Optional[datetime.date], Optional[datetime.date], Optional[datetime.date]]:
    if value is None:
        return None, None, None
    lines = str(value).split("\n")
    lines += [""] * (3 - len(lines))  # 想定外に短い場合の保険
    order_date = _parse_deadline_line(lines[0])
    company_deadline = _parse_deadline_line(lines[1])
    customer_deadline = _parse_deadline_line(lines[2])
    return order_date, company_deadline, customer_deadline


def _parse_stagnation(value) -> tuple[Optional[int], Optional[str]]:
    if value is None or value == "":
        return None, None
    if isinstance(value, (int, float)):
        return int(value), None
    return None, str(value).strip()


def _parse_memo1(value) -> tuple[Optional[str], Optional[datetime.date]]:
    """MEMO1: 8桁のYYYYMMDD数値なら合意納期として、それ以外は自由記述として扱う。"""
    if value is None or value == "":
        return None, None
    if isinstance(value, (int, float)):
        text = str(int(value))
        if len(text) == 8:
            try:
                return None, datetime.date(int(text[0:4]), int(text[4:6]), int(text[6:8]))
            except ValueError:
                pass
        return text, None
    return str(value).strip() or None, None


def _parse_memo2(value) -> Optional[str]:
    if value is None or value == "" or value == 0:
        return None
    return str(value).strip() or None


def _parse_process_step(slot_index: int, value) -> Optional[ProcessStep]:
    if value is None or value == "":
        return None
    lines = str(value).split("\n")
    if len(lines) < 5:
        # 想定外のフォーマット。集計を止めないよう生値をステータスに残して警告対象にする。
        return ProcessStep(
            slot_index=slot_index,
            process_code=lines[0] if lines else "",
            destination=None,
            date1=None,
            date2=None,
            status=f"不明な形式: {value!r}",
        )
    process_code, destination, date1, date2, status = lines[0], lines[1], lines[2], lines[3], lines[-1]
    return ProcessStep(
        slot_index=slot_index,
        process_code=process_code.strip(),
        destination=destination.strip() or None,
        date1=date1.strip() or None,
        date2=date2.strip() or None,
        status=status.strip(),
    )


def load_orders(path: Path) -> list[OrderRecord]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if row and row[0] == HEADER_MARKER:
            header_row_idx = i
            break
    if header_row_idx is None:
        raise ValueError(f"ヘッダー行(先頭列が「{HEADER_MARKER}」の行)が見つかりません。フォーマットをご確認ください。")

    orders: list[OrderRecord] = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        order_no = _clean_str(row[COL_ORDER_NO])
        if order_no is None:
            continue  # 末尾の空行など

        drawing_no, product_name = _split_drawing_product(row[COL_DRAWING_PRODUCT])
        order_date, company_deadline, customer_deadline = _parse_deadline_triplet(row[COL_DEADLINE_RAW])
        stagnation_days, stagnation_note = _parse_stagnation(row[COL_STAGNATION])
        sales_note, sales_agreed_deadline = _parse_memo1(row[COL_MEMO1])

        material_arrived = bool(_clean_str(row[PROCESS_SLOT_START_COL]))
        processes: list[ProcessStep] = []
        for slot in range(1, PROCESS_SLOT_COUNT):  # スロット2〜20(スロット1は材料入荷状況で別扱い)
            col = PROCESS_SLOT_START_COL + slot
            step = _parse_process_step(slot + 1, row[col] if col < len(row) else None)
            if step is not None:
                processes.append(step)

        orders.append(
            OrderRecord(
                order_no=order_no,
                drawing_no=drawing_no,
                product_name=product_name,
                qty=_parse_qty(row[COL_QTY]),
                order_date=order_date,
                company_deadline=company_deadline,
                customer_deadline=customer_deadline,
                remaining_business_days=_parse_remaining_business_days(row[COL_REMAINING_DAYS]),
                customer_order_no=_clean_str(row[COL_CUSTOMER_ORDER_NO]),
                customer_no=_clean_str(row[COL_CUSTOMER_NO]),
                stagnation_days=stagnation_days,
                stagnation_note=stagnation_note,
                old_system_no=_clean_str(row[COL_OLD_SYSTEM_NO]),
                sales_note=sales_note,
                sales_agreed_deadline=sales_agreed_deadline,
                production_note=_parse_memo2(row[COL_MEMO2]),
                shipment_plan_bucket=_clean_str(row[COL_MEMO3]),
                material_arrived=material_arrived,
                processes=processes,
                is_forecast_order=(order_date is None and customer_deadline is None),
            )
        )

    return orders
